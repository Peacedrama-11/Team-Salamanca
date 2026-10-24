"""Resource-bounded, offline CV text parser. Bytes arrive on stdin, never disk."""
import io
import json
import logging
import resource
import sys
import zipfile

MAX_FILE = 8 * 1024 * 1024
MAX_TEXT = 250_000


class DocumentError(ValueError):
    pass


def parse_document(data, kind):
    if not data or len(data) > MAX_FILE:
        raise DocumentError("Choose a nonempty PDF or DOCX smaller than 8 MB.")
    if kind == "pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentError("This file is not a valid PDF. Check its file type.")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentError("Password-protected PDFs cannot be read. Use an unlocked sample or enter skills manually.")
        if len(reader.pages) > 50:
            raise DocumentError("Use a CV with at most 50 pages, or enter skills manually.")
        parts, total = [], 0
        for page in reader.pages:
            # Guard decoded content before pypdf's text-extraction allocations.
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > 8 * 1024 * 1024:
                raise DocumentError("This PDF page is too complex to read safely.")
            part = page.extract_text() or ""
            total += len(part)
            if total > MAX_TEXT:
                raise DocumentError("This document contains too much text. Use a shorter CV.")
            parts.append(part)
        text = "\n".join(parts)
    elif kind == "docx":
        if not data.startswith(b"PK\x03\x04"):
            raise DocumentError("This file is not a valid DOCX. Check its file type.")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            files = archive.infolist()
            if len(files) > 300 or sum(f.file_size for f in files) > 16 * 1024 * 1024:
                raise DocumentError("This DOCX exceeds safe archive limits. Use a smaller document.")
            if len({f.filename for f in files}) != len(files):
                raise DocumentError("This DOCX contains duplicate archive entries.")
            names = set(archive.namelist())
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise DocumentError("This archive is not a DOCX document.")
            for entry in files:
                if entry.flag_bits & 1 or entry.file_size / max(entry.compress_size, 1) > 200:
                    raise DocumentError("This DOCX is encrypted or exceeds safe compression limits.")
                if entry.filename.endswith(".xml"):
                    xml = archive.read(entry).upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise DocumentError("Unsupported XML content in this DOCX.")
            if b"MACROENABLED" in archive.read("[Content_Types].xml").upper():
                raise DocumentError("Macro-enabled documents are not supported.")
        from docx import Document
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        text = "\n".join(parts)
    else:
        raise DocumentError("Choose a PDF or DOCX file.")
    if len(text) > MAX_TEXT:
        raise DocumentError("This document contains too much text. Use a shorter CV.")
    if len("".join(c for c in text if c.isalnum())) < 15:
        raise DocumentError("No readable text found. Scanned/image-only PDFs need OCR, which is not supported. Enter skills manually or use a text-based CV.")
    return text


def main():
    logging.disable(logging.CRITICAL)
    # Linux process-local limits; also bounded and killed by the Node parent.
    resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024,) * 2)
    resource.setrlimit(resource.RLIMIT_CPU, (8, 8))
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    try:
        text = parse_document(sys.stdin.buffer.read(MAX_FILE + 1), sys.argv[1])
        print(json.dumps({"text": text}))
    except DocumentError as error:
        print(json.dumps({"error": str(error)}))
    except Exception:
        print(json.dumps({"error": "This document could not be read safely. Try another PDF/DOCX or enter skills manually."}))


if __name__ == "__main__":
    main()
